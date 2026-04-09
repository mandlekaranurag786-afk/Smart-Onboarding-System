"""
Excel-backed scheduling service utilities.
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path
import logging
import random
import string
from typing import Any, Dict, List, Optional

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

logger = logging.getLogger(__name__)

EXCEL_FILENAME = "updated_scheduling_data.xlsx"
AVAILABILITY_SHEET = "Availability"
SCHEDULED_MEETINGS_SHEET = "Scheduled_Meetings"


def _resolve_excel_path() -> Path:
    app_dir = Path(__file__).resolve().parents[1]  # backend/app
    backend_dir = Path(__file__).resolve().parents[2]  # backend
    project_dir = backend_dir.parent  # Smart-Onboarding-System

    candidates = [
        app_dir / "data" / EXCEL_FILENAME,
        backend_dir / "app" / "data" / EXCEL_FILENAME,
        backend_dir / "data" / EXCEL_FILENAME,
        project_dir / "backend" / "app" / "data" / EXCEL_FILENAME,
    ]
    for file_path in candidates:
        if file_path.exists():
            return file_path
    raise FileNotFoundError(
        f"Could not find {EXCEL_FILENAME}. Tried: {', '.join(str(p) for p in candidates)}"
    )


def _normalize(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _normalized_header(value: Any) -> str:
    return _normalize(value).lower()


def _sheet_by_name(
    workbook: Workbook,
    expected_name: str,
    *,
    create_if_missing: bool = False,
) -> Worksheet:
    expected = expected_name.strip().lower()
    for sheet_name in workbook.sheetnames:
        if sheet_name.strip().lower() == expected:
            return workbook[sheet_name]

    if create_if_missing:
        return workbook.create_sheet(expected_name)

    raise ValueError(f"Sheet '{expected_name}' not found in workbook.")


def _headers_map(sheet: Worksheet) -> Dict[str, int]:
    headers: Dict[str, int] = {}
    for col_idx, cell in enumerate(sheet[1], start=1):
        header = _normalized_header(cell.value)
        if header:
            headers[header] = col_idx
    return headers


def _row_value(sheet: Worksheet, row_idx: int, headers: Dict[str, int], key: str) -> str:
    col_idx = headers.get(key.lower())
    if not col_idx:
        return ""
    return _normalize(sheet.cell(row=row_idx, column=col_idx).value)


def _ensure_meetings_headers(sheet: Worksheet) -> Dict[str, int]:
    expected_headers = [
        "candidate_name",
        "meeting_type",
        "interviewer_name",
        "interviewer_role",
        "date",
        "time",
        "status",
        "booked_by",
        "booked_at",
        "meeting_link",
    ]
    headers = _headers_map(sheet)

    if not headers:
        sheet.append(expected_headers)
        return {header: idx for idx, header in enumerate(expected_headers, start=1)}

    missing = [header for header in expected_headers if header not in headers]
    if missing:
        raise ValueError(
            f"'{SCHEDULED_MEETINGS_SHEET}' sheet is missing required headers: {missing}"
        )
    return headers


def _generate_meeting_link() -> str:
    chunk_1 = "".join(random.choices(string.ascii_lowercase, k=3))
    chunk_2 = "".join(random.choices(string.ascii_lowercase, k=4))
    return f"https://meet.google.com/onboardiq-{chunk_1}-{chunk_2}"


def get_all_available_slots(person_name: str) -> List[Dict[str, str]]:
    workbook = load_workbook(_resolve_excel_path())
    availability = _sheet_by_name(workbook, AVAILABILITY_SHEET)
    headers = _headers_map(availability)

    target_name = person_name.strip().lower()
    slots: List[Dict[str, str]] = []

    for row_idx in range(2, availability.max_row + 1):
        name = _row_value(availability, row_idx, headers, "name")
        available = _row_value(availability, row_idx, headers, "available").upper()
        if name.lower() == target_name and available == "YES":
            slots.append(
                {
                    "name": name,
                    "role": _row_value(availability, row_idx, headers, "role"),
                    "date": _row_value(availability, row_idx, headers, "date"),
                    "time": _row_value(availability, row_idx, headers, "time"),
                }
            )
    return slots


def get_first_available_slot(person_name: str) -> Optional[Dict[str, str]]:
    slots = get_all_available_slots(person_name)
    return slots[0] if slots else None


def get_availability_as_text(person_name: str) -> str:
    workbook = load_workbook(_resolve_excel_path())
    availability = _sheet_by_name(workbook, AVAILABILITY_SHEET)
    headers = _headers_map(availability)

    target_name = person_name.strip().lower()
    lines: List[str] = []
    for row_idx in range(2, availability.max_row + 1):
        name = _row_value(availability, row_idx, headers, "name")
        if name.lower() != target_name:
            continue
        role = _row_value(availability, row_idx, headers, "role")
        date = _row_value(availability, row_idx, headers, "date")
        time = _row_value(availability, row_idx, headers, "time")
        available = _row_value(availability, row_idx, headers, "available")
        lines.append(f"{name} | {role} | {date} | {time} | {available}")

    if not lines:
        return f"No availability rows found for {person_name}"
    return "\n".join(lines)


def get_scheduled_meetings_as_text(candidate_name: str) -> str:
    workbook = load_workbook(_resolve_excel_path())
    try:
        scheduled_sheet = _sheet_by_name(workbook, SCHEDULED_MEETINGS_SHEET)
    except ValueError:
        return "No meetings scheduled yet"

    headers = _headers_map(scheduled_sheet)
    target_name = candidate_name.strip().lower()

    lines: List[str] = []
    for row_idx in range(2, scheduled_sheet.max_row + 1):
        row_candidate_name = _row_value(scheduled_sheet, row_idx, headers, "candidate_name")
        if row_candidate_name.lower() != target_name:
            continue

        meeting_type = _row_value(scheduled_sheet, row_idx, headers, "meeting_type")
        interviewer_name = _row_value(scheduled_sheet, row_idx, headers, "interviewer_name")
        date = _row_value(scheduled_sheet, row_idx, headers, "date")
        time = _row_value(scheduled_sheet, row_idx, headers, "time")
        status = _row_value(scheduled_sheet, row_idx, headers, "status") or "Scheduled"
        lines.append(
            f"{row_candidate_name} | {meeting_type} | {interviewer_name} | {date} | {time} | {status}"
        )

    if not lines:
        return "No meetings scheduled yet"
    return "\n".join(lines)


def ask_llm_for_slot(
    candidate_name: str,
    meeting_type: str,
    primary_person: str,
    backup_person: Optional[str],
    llm,
) -> Dict[str, Any]:
    try:
        logger.info("[LLM SCHEDULER] Asking LLM for slot decision...")

        primary_availability_text = get_availability_as_text(primary_person)
        backup_availability_text = (
            get_availability_as_text(backup_person)
            if backup_person
            else "No backup interviewer provided"
        )
        scheduled_meetings_text = get_scheduled_meetings_as_text(candidate_name)

        prompt = f"""
You are a meeting scheduling assistant for OnboardIQ.

CANDIDATE: {candidate_name}
MEETING TYPE: {meeting_type}
PRIMARY INTERVIEWER: {primary_person}
BACKUP INTERVIEWER: {backup_person or 'None'}

PRIMARY INTERVIEWER AVAILABILITY:
{primary_availability_text}

BACKUP INTERVIEWER AVAILABILITY:
{backup_availability_text}

ALREADY SCHEDULED MEETINGS FOR THIS CANDIDATE:
{scheduled_meetings_text}

RULES:
1. Only pick slots where available column is YES
2. Never pick a slot already in ALREADY SCHEDULED MEETINGS
3. Always prefer primary interviewer over backup
4. Only use backup if primary has zero YES slots
5. If no YES slots exist for anyone return STATUS as PENDING

RESPOND IN EXACTLY THIS FORMAT — no extra text:
STATUS: [SCHEDULED or PENDING]
INTERVIEWER: [exact name as it appears in availability data]
DATE: [exact date as it appears in availability data]
TIME: [exact time as it appears in availability data]
REASON: [one sentence why this slot was chosen]
FALLBACK NOTE: [one sentence about backup situation or N/A]
"""

        llm_response = llm.invoke(prompt)
        raw_response = (
            llm_response.content if hasattr(llm_response, "content") else str(llm_response)
        )
        logger.info("[LLM SCHEDULER] Raw response: %s", raw_response)

        parsed: Dict[str, str] = {}
        for line in raw_response.splitlines():
            if ":" not in line:
                continue
            key, value = line.split(":", 1)
            parsed[key.strip().upper()] = value.strip()

        status = parsed.get("STATUS", "").upper()
        if status == "SCHEDULED":
            return {
                "success": True,
                "status": "SCHEDULED",
                "interviewer": parsed.get("INTERVIEWER", ""),
                "date": parsed.get("DATE", ""),
                "time": parsed.get("TIME", ""),
                "reason": parsed.get("REASON", "LLM selected an available slot."),
                "fallback_note": parsed.get("FALLBACK NOTE", "N/A"),
            }

        return {
            "success": False,
            "status": "PENDING",
            "reason": parsed.get("REASON", "No suitable slot found by LLM."),
            "fallback_note": parsed.get("FALLBACK NOTE", "N/A"),
        }

    except Exception as error:  # noqa: BLE001
        logger.warning("[LLM SCHEDULER ERROR] %s — falling back to deterministic logic", error)

        slot = get_first_available_slot(primary_person)
        fallback_note = "N/A"

        if not slot and backup_person:
            slot = get_first_available_slot(backup_person)
            if slot:
                fallback_note = f"Primary unavailable, selected backup interviewer {backup_person}."

        if slot:
            return {
                "success": True,
                "status": "SCHEDULED",
                "interviewer": slot["name"],
                "date": slot["date"],
                "time": slot["time"],
                "reason": "Fallback to first available slot because LLM was unavailable.",
                "fallback_note": fallback_note,
            }

        return {
            "success": False,
            "status": "PENDING",
            "reason": "No available slots found in fallback scheduling logic.",
            "fallback_note": "N/A",
        }


def get_all_interviewers(role_filter: Optional[str] = None) -> List[Dict[str, str]]:
    workbook = load_workbook(_resolve_excel_path())
    availability = _sheet_by_name(workbook, AVAILABILITY_SHEET)
    headers = _headers_map(availability)

    target_role = role_filter.strip().lower() if role_filter else None
    seen = set()
    interviewers: List[Dict[str, str]] = []

    for row_idx in range(2, availability.max_row + 1):
        name = _row_value(availability, row_idx, headers, "name")
        role = _row_value(availability, row_idx, headers, "role")
        if not name:
            continue
        if target_role and role.lower() != target_role:
            continue
        key = (name.lower(), role.lower())
        if key in seen:
            continue
        seen.add(key)
        interviewers.append({"name": name, "role": role})

    interviewers.sort(key=lambda person: (person["role"].lower(), person["name"].lower()))
    return interviewers


def book_slot(
    candidate_name: str,
    meeting_type: str,
    interviewer_name: str,
    date: str,
    time: str,
    booked_by: str,
) -> Dict[str, Any]:
    workbook = load_workbook(_resolve_excel_path())
    availability = _sheet_by_name(workbook, AVAILABILITY_SHEET)
    availability_headers = _headers_map(availability)

    target_name = interviewer_name.strip().lower()
    target_date = date.strip()
    target_time = time.strip()

    matched_row = None
    interviewer_role = ""

    for row_idx in range(2, availability.max_row + 1):
        name = _row_value(availability, row_idx, availability_headers, "name")
        slot_date = _row_value(availability, row_idx, availability_headers, "date")
        slot_time = _row_value(availability, row_idx, availability_headers, "time")
        if name.lower() == target_name and slot_date == target_date and slot_time == target_time:
            matched_row = row_idx
            interviewer_role = _row_value(availability, row_idx, availability_headers, "role")
            break

    if matched_row is None:
        return {
            "success": False,
            "error": "slot_not_found",
            "message": "Slot not found for the selected interviewer/date/time.",
        }

    available_col = availability_headers.get("available")
    if not available_col:
        raise ValueError(f"'{AVAILABILITY_SHEET}' sheet is missing the 'available' column.")

    current_status = _normalize(
        availability.cell(row=matched_row, column=available_col).value
    ).upper()
    if current_status != "YES":
        return {
            "success": False,
            "error": "slot_unavailable",
            "message": "Selected slot is already booked.",
        }

    availability.cell(row=matched_row, column=available_col, value="NO")

    scheduled_sheet = _sheet_by_name(
        workbook,
        SCHEDULED_MEETINGS_SHEET,
        create_if_missing=True,
    )
    scheduled_headers = _ensure_meetings_headers(scheduled_sheet)

    booking = {
        "candidate_name": candidate_name,
        "meeting_type": meeting_type,
        "interviewer_name": interviewer_name,
        "interviewer_role": interviewer_role,
        "date": date,
        "time": time,
        "status": "confirmed",
        "booked_by": booked_by,
        "booked_at": datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC"),
        "meeting_link": _generate_meeting_link(),
    }

    row_values = [booking[header] for header in sorted(scheduled_headers, key=scheduled_headers.get)]
    scheduled_sheet.append(row_values)
    workbook.save(_resolve_excel_path())

    return {"success": True, "booking": booking}


def get_meetings_for_candidate(candidate_name: str) -> List[Dict[str, str]]:
    workbook = load_workbook(_resolve_excel_path())
    try:
        scheduled_sheet = _sheet_by_name(workbook, SCHEDULED_MEETINGS_SHEET)
    except ValueError:
        return []

    headers = _headers_map(scheduled_sheet)
    normalized_candidate_name = candidate_name.strip().lower()
    meetings: List[Dict[str, str]] = []
    for row_idx in range(2, scheduled_sheet.max_row + 1):
        sheet_candidate_name = _row_value(scheduled_sheet, row_idx, headers, "candidate_name")
        if sheet_candidate_name.lower() != normalized_candidate_name:
            continue
        meetings.append(
            {
                "candidate_name": sheet_candidate_name,
                "meeting_type": _row_value(scheduled_sheet, row_idx, headers, "meeting_type"),
                "interviewer_name": _row_value(
                    scheduled_sheet, row_idx, headers, "interviewer_name"
                ),
                "interviewer_role": _row_value(
                    scheduled_sheet, row_idx, headers, "interviewer_role"
                ),
                "date": _row_value(scheduled_sheet, row_idx, headers, "date"),
                "time": _row_value(scheduled_sheet, row_idx, headers, "time"),
                "status": _row_value(scheduled_sheet, row_idx, headers, "status"),
                "booked_by": _row_value(scheduled_sheet, row_idx, headers, "booked_by"),
                "booked_at": _row_value(scheduled_sheet, row_idx, headers, "booked_at"),
                "meeting_link": _row_value(scheduled_sheet, row_idx, headers, "meeting_link"),
            }
        )
    return meetings


def get_interviewer_email(interviewer_name: str) -> str | None:
    """
    Reads the Availability sheet and returns the mail id
    for the given interviewer name.
    Returns the first match found.
    Returns None if not found.
    """
    target_name = _normalize(interviewer_name)
    if not target_name:
        return None

    excel_path = Path(__file__).resolve().parents[1] / "data" / "updated_scheduling_data.xlsx"
    if not excel_path.exists():
        logger.warning("[SCHEDULING] Availability file not found for interviewer email lookup: %s", excel_path)
        return None

    try:
        workbook = load_workbook(excel_path)
        availability = workbook[AVAILABILITY_SHEET]
    except Exception as exc:  # noqa: BLE001
        logger.warning("[SCHEDULING] Failed to read interviewer email from Excel: %s", exc)
        return None

    for row in availability.iter_rows(min_row=2):
        if len(row) < 6:
            continue

        row_name = _normalize(row[0].value)
        row_email = _normalize(row[5].value)
        if row_name and row_name.lower() == target_name.lower():
            return row_email or None

    return None
