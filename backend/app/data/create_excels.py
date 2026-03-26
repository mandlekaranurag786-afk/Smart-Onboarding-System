from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from datetime import datetime, timedelta

# Create workbook
wb = Workbook()
ws = wb.active
ws.title = "availability"

# Define headers
headers = ["hr_name", "role", "date", "time_slot", "status", "candidate_name"]

# Add headers
ws.append(headers)

# Style headers
for col in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col)
    cell.font = Font(bold=True)
    cell.alignment = Alignment(horizontal="center")

# HR data (aligned with your UI)
hr_data = [
    {"name": "Mohini Moghe", "role": "Senior HR Manager"},
    {"name": "Rahul Sharma", "role": "HR Manager"}
]

# Time slots (exactly like UI)
time_slots = ["09:00 AM", "11:30 AM", "02:00 PM", "04:30 PM"]

# Generate for next 7 days
start_date = datetime.today()

row_num = 2

for day in range(7):
    current_date = (start_date + timedelta(days=day)).strftime("%Y-%m-%d")

    for hr in hr_data:
        for slot in time_slots:
            ws.cell(row=row_num, column=1, value=hr["name"])
            ws.cell(row=row_num, column=2, value=hr["role"])
            ws.cell(row=row_num, column=3, value=current_date)
            ws.cell(row=row_num, column=4, value=slot)
            ws.cell(row=row_num, column=5, value="available")  # default
            ws.cell(row=row_num, column=6, value="")  # candidate empty

            row_num += 1

# Auto-adjust column width
for col in ws.columns:
    max_length = 0
    col_letter = col[0].column_letter

    for cell in col:
        if cell.value:
            max_length = max(max_length, len(str(cell.value)))

    ws.column_dimensions[col_letter].width = max_length + 2

# Save file
file_name = "hr_meeting_availability.xlsx"
wb.save(file_name)

print(f"✅ Excel '{file_name}' generated successfully")